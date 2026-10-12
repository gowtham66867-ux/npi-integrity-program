"""L1 ingestion & landing: source registry, write-once raw zone, validation, staging.

Layout of the inbox and raw zone (same shape):
    <zone>/<source_key>/<version YYYY-MM-DD>/<original file name>

Files are copied into the raw zone exactly as received, stamped with a SHA-256,
and made read-only. A file with the same path but different bytes is refused:
a new release must arrive as a new version folder.
"""
from __future__ import annotations

import csv
import hashlib
import os
import shutil
import stat
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import duckdb

from . import config


@dataclass(frozen=True)
class Source:
    key: str
    label: str                 # used in evidence citations
    kind: str                  # "snapshot" = latest version <= as-of; "cumulative" = all versions <= as-of
    required: tuple[str, ...]  # columns that must exist
    key_column: str            # column checked for nulls
    columns: tuple[str, ...] | None = None  # projection (None = all columns)
    phase: str = "Prototype"
    origin: str = ""


NPPES_KEEP = (
    "NPI", "Entity Type Code", "Provider Organization Name (Legal Business Name)", "Provider Last Name (Legal Name)",
    "Provider First Name", "Provider Middle Name", "Provider Credential Text",
    "Provider First Line Business Practice Location Address",
    "Provider Second Line Business Practice Location Address",
    "Provider Business Practice Location Address City Name", "Provider Business Practice Location Address State Name",
    "Provider Business Practice Location Address Postal Code",
    "Provider Business Practice Location Address Telephone Number", "Provider Enumeration Date", "Last Update Date",
    "NPI Deactivation Date", "NPI Reactivation Date", "Authorized Official Last Name",
    "Authorized Official First Name", "Authorized Official Title or Position", "Healthcare Provider Taxonomy Code_1",
)

SOURCES = {s.key: s for s in [
    Source("nppes_full", "NPPES full monthly file", "snapshot", NPPES_KEEP, "NPI", NPPES_KEEP,
           origin="https://download.cms.gov/nppes/NPI_Files.html"),
    Source("nppes_weekly", "NPPES weekly incremental file", "cumulative", NPPES_KEEP, "NPI", NPPES_KEEP,
           origin="https://download.cms.gov/nppes/NPI_Files.html"),
    Source("nppes_practice_locations", "NPPES practice location reference file", "snapshot", ("NPI",), "NPI",
           origin="https://download.cms.gov/nppes/NPI_Files.html"),
    Source("nppes_deactivation", "NPPES Deactivated NPI Report", "cumulative", ("NPI", "NPPES Deactivation Date"),
           "NPI", origin="https://download.cms.gov/nppes/NPI_Files.html"),
    Source("oig_leie", "OIG LEIE (UPDATED.csv)", "snapshot", ("LASTNAME", "FIRSTNAME", "NPI", "EXCLDATE"),
           "EXCLDATE", origin="https://oig.hhs.gov/exclusions/exclusions_list.asp"),
    Source("sam_exclusions", "SAM.gov exclusions extract", "snapshot", ("First", "Last", "Active Date"),
           "Active Date", origin="https://open.gsa.gov/api/"),
    Source("pecos_enrollment", "PECOS Public Provider Enrollment", "snapshot", ("NPI", "ENRLMT_ID"), "NPI",
           origin="https://data.cms.gov"),
    Source("pecos_reassignment", "PECOS reassignment of benefits", "snapshot",
           ("Group Enrollment ID", "Individual NPI"), "Individual NPI", origin="https://data.cms.gov"),
    Source("order_referring", "Medicare Order and Referring file", "snapshot", ("NPI", "PARTB", "DME", "HHA"), "NPI",
           origin="https://data.cms.gov"),
    Source("optout", "Medicare opt-out affidavits", "snapshot",
           ("NPI", "Optout Effective Date", "Optout End Date"), "NPI", origin="https://data.cms.gov"),
    Source("medicare_by_provider", "Medicare Physician & Other Practitioners by Provider", "snapshot",
           ("Rndrng_NPI", "Tot_Srvcs", "Tot_Benes"), "Rndrng_NPI", origin="https://data.cms.gov"),
    Source("claims", "Claim lines (synthetic 837 feed)", "cumulative",
           ("LINE_ID", "SRVC_DT", "BILLING_NPI", "RENDERING_NPI", "HCPCS_CD"), "LINE_ID",
           origin="synthetic (prototype); 837/835 under data-use agreement in production"),
    Source("zip_centroids", "ZIP centroid reference", "snapshot", ("ZIP", "LAT", "LON"), "ZIP",
           origin="synthetic stand-in for the Census ZCTA gazetteer"),
    Source("address_types", "Address type reference (virtual office / mail drop)", "snapshot",
           ("ADDRESS_LINE_1", "ZIP", "ADDRESS_TYPE"), "ZIP",
           origin="synthetic stand-in for a USPS CMRA / address-validation feed"),
]}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


class WriteOnceViolation(RuntimeError):
    pass


def land(inbox: Path, raw: Path = config.RAW) -> list[dict]:
    """Copy every inbox file into the raw zone (write-once). Returns the landed-file manifest."""
    landed = []
    for src_dir in sorted(p for p in inbox.iterdir() if p.is_dir()):
        if src_dir.name not in SOURCES:
            continue
        for ver_dir in sorted(p for p in src_dir.iterdir() if p.is_dir()):
            for f in sorted(p for p in ver_dir.iterdir() if p.is_file()):
                dest = raw / src_dir.name / ver_dir.name / f.name
                digest = sha256(f)
                if dest.exists():
                    if sha256(dest) != digest:
                        raise WriteOnceViolation(
                            f"{dest} already exists with different content; publish a new version folder instead")
                    status = "unchanged"
                else:
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(f, dest)
                    os.chmod(dest, stat.S_IRUSR | stat.S_IRGRP | stat.S_IROTH)
                    status = "landed"
                landed.append({"source": src_dir.name, "version": ver_dir.name, "file": f.name,
                               "path": str(dest.relative_to(raw)), "sha256": digest, "status": status})
    return landed


def raw_manifest(raw: Path = config.RAW, as_of: str | None = None) -> list[dict]:
    """Every file in the raw zone (optionally only versions <= as_of), with checksums."""
    out = []
    for src_dir in sorted(p for p in raw.iterdir() if p.is_dir() and p.name in SOURCES):
        for ver_dir in sorted(p for p in src_dir.iterdir() if p.is_dir()):
            if as_of and ver_dir.name > as_of:
                continue
            for f in sorted(p for p in ver_dir.iterdir() if p.is_file()):
                out.append({"source": src_dir.name, "version": ver_dir.name, "file": f.name,
                            "path": str(f.relative_to(raw)), "sha256": sha256(f)})
    return out


def select_files(manifest: list[dict], as_of: str) -> list[dict]:
    """Apply snapshot/cumulative semantics: which files feed a build as of a date."""
    chosen = []
    for key, src in SOURCES.items():
        files = [m for m in manifest if m["source"] == key and m["version"] <= as_of]
        if not files:
            continue
        if src.kind == "snapshot":
            latest = max(m["version"] for m in files)
            files = [m for m in files if m["version"] == latest]
        chosen.extend(files)
    return chosen


def _xlsx_to_csv(path: Path, cache: Path) -> Path:
    """The real deactivation report ships as .xlsx with banner rows above the header."""
    import openpyxl
    out = cache / (path.stem + ".csv")
    if out.exists():
        return out
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = None
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        for r in rows:
            vals = ["" if v is None else (v.strftime("%m/%d/%Y") if hasattr(v, "strftime") else str(v)) for v in r]
            if header is None:
                if any(v.strip().upper() == "NPI" for v in vals):
                    header = [v.strip() for v in vals]
                    w.writerow(header)
                continue
            if any(vals):
                w.writerow(vals)
    return out


def q(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def stage(conn: duckdb.DuckDBPyConnection, files: list[dict], raw: Path = config.RAW) -> list[dict]:
    """Load the selected raw files into stg_<source> tables with lineage columns and validate them."""
    cache = config.DATA / "stage_cache"
    cache.mkdir(exist_ok=True)
    results = []
    by_source: dict[str, list[dict]] = {}
    for m in files:
        by_source.setdefault(m["source"], []).append(m)
    for key, src in SOURCES.items():
        conn.execute(f"DROP TABLE IF EXISTS stg_{key}")
        prev_rows = None
        for m in sorted(by_source.get(key, []), key=lambda x: x["version"]):
            path = raw / m["path"]
            if path.suffix.lower() == ".xlsx":
                path = _xlsx_to_csv(path, cache)
            reader = f"read_csv({str(path)!r}, all_varchar=true, header=true, sample_size=-1)"
            cols = conn.execute(f"DESCRIBE SELECT * FROM {reader}").fetchall()
            present = [c[0] for c in cols]
            missing = [c for c in src.required if c not in present]
            checks = {"missing_columns": missing}
            if missing:
                results.append({**m, "rows": 0, "status": "failed", "checks": checks})
                continue
            keep = src.columns or tuple(present)
            proj = ", ".join(q(c) for c in keep if c in present)
            conn.execute(f"CREATE OR REPLACE TEMP TABLE _load AS SELECT {proj} FROM {reader}")
            exists = conn.execute(
                "SELECT count(*) FROM information_schema.tables WHERE table_name = ?", [f"stg_{key}"]).fetchone()[0]
            lineage = (f"DATE '{m['version']}' AS _version, {m['file']!r} AS _file, "
                       f"{m['sha256'][:16]!r} AS _sha, rowid + 1 AS _row")
            if not exists:
                conn.execute(f"CREATE TABLE stg_{key} AS SELECT *, {lineage} FROM _load")
            else:
                conn.execute(f"INSERT INTO stg_{key} BY NAME SELECT *, {lineage} FROM _load")
            n, nulls = conn.execute(
                f"SELECT count(*), count(*) FILTER (WHERE coalesce(trim({q(src.key_column)}), '') = '') FROM _load"
            ).fetchone()
            checks["null_key_rate"] = round(nulls / n, 4) if n else 0
            status = "ok"
            if src.kind == "snapshot" and prev_rows:
                drift = abs(n - prev_rows) / prev_rows
                checks["row_drift_vs_prior"] = round(drift, 3)
                if drift > 0.25:
                    status = "warning"
            if checks["null_key_rate"] > 0.05:
                status = "warning"
            prev_rows = n
            results.append({**m, "rows": n, "status": status, "checks": checks})
    return results
